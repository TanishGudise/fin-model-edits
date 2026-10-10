"""Parse and apply cell edits: the answer format, and the format task data is written in.

A reply contains one or more ```edits blocks. Each non-blank line sets one cell or one range:

    IS!F14 = =SUM(F11:F13)          formula
    IS!F14:J14 = =SUM(F11:F13)      formula filled across the range, like Excel's fill:
                                    relative references shift, $-anchored ones don't
    Inputs!C23 = 0.23               number (23% is accepted too)
    Inputs!A9 = Revenue growth      text
    CF!F6 =                         clear the cell

Lines starting with `#` are comments. Edits apply in order, so a later edit to the same cell
wins. Any line that doesn't parse, an unknown sheet, or a cell outside the grid rejects the
whole reply: a half-applied answer would be graded on something the model didn't write.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import coordinate_from_string, get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook

MAX_ROW = 200
MAX_COL = 26  # A..Z
MAX_CELLS = 5000
MAX_CONTENT_CHARS = 1000

_BLOCK = re.compile(r"^```edits[ \t]*\n(?P<body>.*?)^```", re.MULTILINE | re.DOTALL)
_LINE = re.compile(
    r"^(?P<sheet>[A-Za-z][A-Za-z0-9_]*)!(?P<ref>[A-Za-z$]+[0-9$]+(?::[A-Za-z$]+[0-9$]+)?)\s*=(?P<content>.*)$"
)
_NUMBER = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$|^[+-]?\.\d+$")

Value = float | int | str | None
"""A cell's new content: a number, a formula (str starting with "="), text, or None (clear)."""


class EditError(ValueError):
    pass


@dataclass(frozen=True)
class Edit:
    sheet: str
    ref: str  # "F14" or "F14:J14", no $ signs
    content: Value

    def cells(self) -> list[tuple[str, Value]]:
        """(coordinate, content) for every cell the edit sets, with formulas filled."""
        min_col, min_row, max_col, max_row = range_boundaries(self.ref)
        origin = f"{get_column_letter(min_col)}{min_row}"
        out = []
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                coord = f"{get_column_letter(col)}{row}"
                content = self.content
                if isinstance(content, str) and content.startswith("=") and coord != origin:
                    content = Translator(content, origin=origin).translate_formula(coord)
                out.append((coord, content))
        return out


def parse_content(raw: str) -> Value:
    text = raw.strip()
    if not text:
        return None
    if text.startswith("="):
        return normalize_formula(text)
    if text.endswith("%") and _NUMBER.match(text[:-1].strip()):
        return _to_number(text[:-1].strip()) / 100
    if _NUMBER.match(text):
        return _to_number(text)
    if len(text) >= 2 and text[0] == text[-1] and text[0] in "\"'":
        text = text[1:-1]
    return text


def _to_number(text: str) -> float | int:
    text = text.replace(",", "")
    value = float(text)
    return int(value) if value.is_integer() and "." not in text else value


def normalize_formula(formula: str) -> str:
    """Upper-case function names (LibreOffice reads `=sum(` as an unknown name)."""
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return formula
    if not any(
        t.type == Token.FUNC and t.subtype == Token.OPEN and t.value != t.value.upper() for t in tokens
    ):
        return formula
    return "=" + "".join(
        t.value.upper() if t.type == Token.FUNC and t.subtype == Token.OPEN else t.value for t in tokens
    )


def parse_edit_lines(text: str, sheets: list[str]) -> list[Edit]:
    """Parse the body of one edits block (or a task data file)."""
    by_lower = {name.lower(): name for name in sheets}
    edits: list[Edit] = []
    for number, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        match = _LINE.match(stripped)
        if not match:
            raise EditError(f"line {number} is not `Sheet!Cell = content`: {stripped[:80]!r}")
        sheet = by_lower.get(match["sheet"].lower())
        if sheet is None:
            raise EditError(f"line {number}: unknown sheet {match['sheet']!r}")
        ref = match["ref"].replace("$", "").upper()
        for part in ref.split(":"):
            col, row = coordinate_from_string(part)
            if not (1 <= row <= MAX_ROW) or len(col) > 1:
                raise EditError(f"line {number}: {part} is outside A1:Z{MAX_ROW}")
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        if min_col > max_col or min_row > max_row:
            raise EditError(f"line {number}: range {ref} is reversed")
        raw = match["content"]
        if raw.startswith(" "):
            raw = raw[1:]
        if len(raw) > MAX_CONTENT_CHARS:
            raise EditError(f"line {number}: content longer than {MAX_CONTENT_CHARS} characters")
        edits.append(Edit(sheet, ref, parse_content(raw)))
    if sum(len(e.cells()) for e in edits) > MAX_CELLS:
        raise EditError(f"more than {MAX_CELLS} cells edited")
    return edits


def find_blocks(reply: str) -> list[str]:
    return [m["body"] for m in _BLOCK.finditer(reply or "")]


def parse_reply(reply: str, sheets: list[str]) -> list[Edit] | None:
    """All edits in the reply's ```edits blocks, in order; None if there is no block."""
    blocks = find_blocks(reply)
    if not blocks:
        return None
    return [edit for body in blocks for edit in parse_edit_lines(body, sheets)]


def apply_edits(wb: Workbook, edits: list[Edit]) -> dict[tuple[str, str], Value]:
    """Apply edits in place; returns the final content of every edited cell."""
    changed: dict[tuple[str, str], Value] = {}
    for edit in edits:
        ws = wb[edit.sheet]
        for coord, content in edit.cells():
            ws[coord].value = content
            changed[(edit.sheet, coord)] = content
    return changed


def format_content(value: Value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return repr(value)
    return str(value)


def format_reply(lines: list[str], preamble: str = "") -> str:
    body = "\n".join(lines)
    return (preamble + "\n\n" if preamble else "") + f"```edits\n{body}\n```\n"
