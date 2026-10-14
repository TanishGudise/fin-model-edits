"""The document-reconciliation tasks (fin_model_edits/recon) and what they need: comparisons and IF in
the expression language, the EDGAR HTML conversion with blackline markers, the source manifest, and the
Hillman tasks' build checks. Tests that need the downloaded filings are skipped without them
(scripts/fetch_sources.py hillman); scripts/validate_recon.py grades the tasks in LibreOffice."""

import pytest

from fin_model_edits.gen.expr import Cmp, Fn, In, Ref
from fin_model_edits.gen.layout import Compiler, Layout
from fin_model_edits.gen.model import Evaluator, Line, Model
from fin_model_edits.sources.convert import html_to_text
from fin_model_edits.sources.fetch import check, manifest


def _tiny() -> tuple[Model, Layout]:
    m = Model(company="T", sector="x", units="$ in millions", periods=["Y1", "Y2"], n_actual=0)
    m.add(Line("hi", "hi", "inputs", kind="scalar_input", cells={"": 3.0}))
    m.add(Line("lev", "lev", "debt", cells={"Y1": 3.5, "Y2": 2.0}))
    tier = Fn("IF", Cmp(">", Ref("lev"), In("hi")), 0.5, Fn("IF", Cmp(">", Ref("lev"), 2.5), 0.25, 0))
    m.add(Line("pct", "pct", "debt", cells={"Y1": tier, "Y2": tier}))
    lay = Layout(
        family="R",
        sheets=["S"],
        names={"S": "S"},
        cols={"Y1": "C", "Y2": "D"},
        rows={"lev": ("S", 2), "pct": ("S", 3)},
        scalars={"hi": ("S", "B1")},
    )
    return m, lay


def test_if_with_comparisons_evaluates_and_compiles():
    m, lay = _tiny()
    ev = Evaluator(m)
    assert ev.value("pct", "Y1") == 0.5 and ev.value("pct", "Y2") == 0
    assert Compiler(m, lay).formula("pct", "Y2") == "=IF(D2>$B$1,0.5,IF(D2>2.5,0.25,0))"


def test_if_evaluates_only_the_branch_taken():
    m, _ = _tiny()
    m.add(Line("zero", "zero", "debt", cells={"Y1": 0.0, "Y2": 0.0}))
    m.lines["pct"].cells["Y1"] = Fn("IF", Cmp(">", Ref("lev"), 1), 1, Ref("lev") / Ref("zero"))
    assert Evaluator(m).value("pct", "Y1") == 1


def test_comparisons_only_inside_if():
    from fin_model_edits.gen.model import EvalError

    m, lay = _tiny()
    m.lines["pct"].cells["Y1"] = Cmp(">", Ref("lev"), 1)
    with pytest.raises(EvalError):
        Evaluator(m).value("pct", "Y1")
    with pytest.raises(TypeError):
        Compiler(m, lay).formula("pct", "Y1")
    m.lines["pct"].cells["Y1"] = Fn("IF", Cmp(">", Ref("lev"), 1), 1)
    with pytest.raises(EvalError):
        Evaluator(m).value("pct", "Y1")


def test_blackline_markers_only_where_asked():
    html = (
        "<p>Pay <span style='color:#0000ff;text-decoration:underline'>0.25% of the total</span>"
        "<span style='color:#ff0000;text-decoration:line-through'>1.00%</span> quarterly, a "
        "<span style='color:#000000;text-decoration:underline'>Defined Term</span>.</p>"
        "<table><tr><td>Net sales</td><td>1,695</td></tr></table>"
    )
    marked = html_to_text(html, blackline=True)
    assert "[inserted: 0.25% of the total]" in marked and "[deleted: 1.00%]" in marked
    assert "[inserted: Defined Term]" not in marked and "Defined Term" in marked
    assert "Net sales | 1,695" in marked
    plain = html_to_text(html)
    assert "[inserted" not in plain and "0.25% of the total" in plain


def test_older_blackline_styles():
    html = (
        "<p>Rate <strike>3.00%</strike><font color='#0000FF'>2.50%</font> and "
        "<span style='border-bottom: 1px double #000'>New Term Loans</span> plus <s>old</s>.</p>"
    )
    assert html_to_text(html, blackline=True).strip() == (
        "Rate [deleted: 3.00%][inserted: 2.50%] and [inserted: New Term Loans] plus [deleted: old]."
    )
    assert html_to_text(html).strip() == "Rate 3.00%2.50% and New Term Loans plus old."


def test_manifest_records_each_filing():
    docs = manifest("hillman")
    assert len({d.id for d in docs}) == len(docs)
    for d in docs:
        assert d.url.startswith("https://www.sec.gov/Archives/edgar/data/1822492/")
        assert d.accession.replace("-", "") in d.url and len(d.sha256) == 64
    assert [d.id for d in docs if d.blackline] == ["amendment_no_1"]


def test_amneal_manifest_ships_the_amendments_as_filed():
    docs = manifest("amneal")
    assert len({d.id for d in docs}) == len(docs) == 26
    for d in docs:
        assert d.url.startswith("https://www.sec.gov/Archives/edgar/data/1723128/")
        assert d.accession.replace("-", "") in d.url and len(d.sha256) == 64
    ids = [d.id for d in docs]
    for amendment in ("amendment_no_1", "amendment_no_2", "amendment_no_3", "amendment_no_4"):
        assert amendment in ids


needs_sources = pytest.mark.skipif(
    not all(check(d) for d in manifest("hillman")), reason="Hillman filings not downloaded"
)


@needs_sources
def test_tasks_build_with_their_evidence_in_the_documents():
    from fin_model_edits.recon import checks
    from fin_model_edits.recon.issues import DECOYS, ISSUES, TASKS
    from fin_model_edits.recon.make import make, notes_text
    from fin_model_edits.recon.packet import build

    parts = build()
    for task_id, ids in TASKS.items():
        g = make(task_id)
        assert [i.kind for i in g.issues] == ids
        chains = {i: ISSUES[i].chain for i in ids} | {k: d.evidence for k, d in DECOYS.items()}
        assert checks.evidence_problems(parts, notes_text(g), chains) == []
    control = checks.search_baseline(parts, ISSUES["term_loan_margin"])
    assert control["found_by_search"]
    assert not checks.search_baseline(parts, ISSUES["leverage_gross_of_cash"])["found_by_search"]


def test_amneal_actuals_tie_to_the_10k():
    from fin_model_edits.gen.generate import _values
    from fin_model_edits.recon.amneal import build_clean

    v = _values(build_clean())
    assert round(v[("ni", "FY2025A")], 3) == 127.933  # reported net income
    assert round(v[("ebitda", "FY2025A")], 3) == 603.872  # the release's FY2025 EBITDA
    assert round(v[("ta", "FY2025A")], 2) == 3678.28  # total assets, December 31, 2025
    assert abs(v[("check", "FY2025A")]) < 1e-6


needs_amneal = pytest.mark.skipif(
    not all(check(d) for d in manifest("amneal")), reason="Amneal filings not downloaded"
)


@needs_amneal
def test_amneal_evidence_is_in_the_full_filings():
    from fin_model_edits.recon import checks
    from fin_model_edits.recon.companies import company
    from fin_model_edits.recon.make import make, notes_text

    co = company("amneal")
    parts = co.packet()
    for task_id, ids in co.tasks.items():
        g = make(task_id)
        chains = {i: co.issues[i].chain for i in ids} | {k: d.evidence for k, d in co.decoys.items()}
        assert checks.evidence_problems(parts, notes_text(g), chains) == []
    assert checks.search_baseline(parts, co.issues["tlb_margin"])["found_by_search"]


@needs_amneal
def test_new_structure_is_absent_from_the_task_and_free_in_the_reference(tmp_path):
    import openpyxl

    from fin_model_edits.recon.companies import company
    from fin_model_edits.recon.make import make, write
    from fin_model_edits.tasks import load_spec

    co = company("amneal")
    task_id = next(t for t, ids in co.tasks.items() if "tlb_addon_missing" in ids)
    g = make(task_id)
    absent = co.issues["tlb_addon_missing"].absent
    assert all(key not in g.model.lines and key in g.ref.lines for key in absent)
    spec = load_spec(write(g, tmp_path))
    cells = {g.layout.addr(g.ref, key, p) for key in absent for p in g.ref.lines[key].cells}
    assert {tuple(f.split("!")) for f in spec.free} == cells
    wb = openpyxl.load_workbook(spec.task_xlsx)
    for sheet, coord in cells:
        row = "".join(ch for ch in coord if ch.isdigit())
        assert wb[sheet][coord].value is None and wb[sheet][f"A{row}"].value is None
    assert any(g.ref.lines[k].label in spec.reference for k in absent)  # the reference adds the labels


def test_orphan_rows_flag_a_line_the_task_leaves_unread():
    import openpyxl

    from fin_model_edits.recon.checks import orphan_rows

    def wb(formula):
        w = openpyxl.Workbook()
        ws = w.active
        ws.title = "S"
        ws["A1"], ws["B1"] = "Adj. EBITDA", 10
        ws["A2"], ws["B2"] = "Credit EBITDA", "=B1+1"
        ws["A3"], ws["B3"] = "Leverage", formula
        return w

    assert orphan_rows(wb("=5/B1"), wb("=5/B2")) == ["S!2 'Credit EBITDA'"]
    assert orphan_rows(wb("=5/B2"), wb("=5/B2")) == []


def test_collapse_renders_source_line_breaks_as_spaces():
    html = "<p>the Amendment No. 3\nTerm Loans <span style='color:#0000ff'>outstanding\n on</span> it</p>"
    assert "No. 3\nTerm" in html_to_text(html)
    assert html_to_text(html, blackline=True, collapse=True).strip() == (
        "the Amendment No. 3 Term Loans [inserted: outstanding on] it"
    )


def test_moving_a_built_line_rewrites_references_to_it():
    from fin_model_edits.recon.make import _replace_ref

    assert _replace_ref("=$C$5+$C$55", "Debt", "C55", "C105", "Debt") == "=$C$5+$C$105"
    assert _replace_ref("=Debt!C55*2", "Debt", "C55", "C105", "Financials") == "=Debt!C105*2"
    assert _replace_ref("=C550+C55", "Debt", "C55", "C105", "Debt") == "=C550+C105"


@needs_amneal
def test_judgement_tasks_carry_notes_extra_lines_and_a_leave_alone_task(tmp_path):
    import json

    import openpyxl

    from fin_model_edits.recon.make import make, write
    from fin_model_edits.tasks import load_spec

    a = load_spec(write(make("amneal-j-a"), tmp_path))
    meta = json.loads((a.folder / "meta.json").read_text())
    notes = meta["analyst_notes"]
    assert notes["net1l"]["planted"] and not notes["syn_a"]["planted"]  # notes sit on right rows too
    assert {i["category"] for i in meta["issues"]} == {"wrong method"}
    ws = openpyxl.load_workbook(a.task_xlsx)["Debt"]
    assert any(str(c.value).startswith("Gross of cash") for c in ws["J"] if c.value)
    assert a.free  # the extra add-back row is free in the reference: a reply may clear or reuse it

    b = make("amneal-j-b")
    assert "ecf_ad" not in b.model.lines and "ecf_ad" not in b.ref.lines  # only task (a) has the extra row

    c = load_spec(write(make("amneal-j-c"), tmp_path))
    assert not c.issues and c.reference.strip().splitlines()[1:] == []  # nothing to fix
    assert len(c.decoys) >= 8
