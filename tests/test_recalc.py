"""LibreOffice must recompute every formula, not trust values cached in the file."""

import re
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pytest

from fin_model_edits.recalc import LibreOfficeMissing, _run_soffice, read_values, recalc_sync, soffice_path


def _has_libreoffice() -> bool:
    try:
        soffice_path()
    except LibreOfficeMissing:
        return False
    return True


pytestmark = pytest.mark.skipif(not _has_libreoffice(), reason="LibreOffice is not installed")


def _with_stale_cache(tmp: Path) -> Path:
    """A workbook where A2 = A1*2 = 20, but the file's cached value for A2 says 999."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 10
    wb.active["A2"] = "=A1*2"
    fresh = recalc_sync([wb])
    assert fresh.values is not None and fresh.values[0]["Sheet"]["A2"] == 20
    source = tmp / "source.xlsx"
    wb.save(source)
    # Let LibreOffice write the file once, so it carries cached values, then corrupt one.
    out = tmp / "cached"
    out.mkdir()
    profile = tmp / "first-profile"
    ok, log, _ = _run_soffice(profile, [source], out, tmp)
    assert ok, log
    stale = tmp / "stale.xlsx"
    with zipfile.ZipFile(out / "source.xlsx") as zin, zipfile.ZipFile(stale, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text, n = re.subn(
                    r'(<c r="A2"[^>]*><f[^>]*>[^<]*</f><v>)[^<]*(</v>)', r"\g<1>999\g<2>", data.decode()
                )
                assert n == 1
                data = text.encode()
            zout.writestr(item, data)
    assert openpyxl.load_workbook(stale, data_only=True).active["A2"].value == 999
    return stale


def test_stale_cached_values_are_recalculated():
    with tempfile.TemporaryDirectory() as tmp:
        stale = _with_stale_cache(Path(tmp))
        result = recalc_sync([openpyxl.load_workbook(stale)])
        assert result.values is not None and result.values[0]["Sheet"]["A2"] == 20

        # The same file through a profile without the recalc setting keeps the stale value.
        # That is why every grade uses a profile with "recalculate on load: always".
        out = Path(tmp) / "default-out"
        out.mkdir()
        ok, log, _ = _run_soffice(Path(tmp) / "default-profile", [stale], out, Path(tmp))
        assert ok, log
        assert read_values(out / "stale.xlsx")["Sheet"]["A2"] == 999


def test_parallel_recalcs_do_not_interfere():
    from concurrent.futures import ThreadPoolExecutor

    def one(i: int) -> object:
        wb = openpyxl.Workbook()
        wb.active["A1"] = i
        wb.active["A2"] = "=A1*A1"
        result = recalc_sync([wb])
        return result.values[0]["Sheet"]["A2"] if result.values else result.error

    with ThreadPoolExecutor(8) as pool:
        assert list(pool.map(one, range(16))) == [i * i for i in range(16)]
